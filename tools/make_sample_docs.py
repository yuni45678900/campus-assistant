# -*- coding: utf-8 -*-
"""
生成一组演示用的校园文档（虚构学校"明德大学"），用于在真实资料就位前
验证整条链路是否正常。

    python tools/make_sample_docs.py

产物写入 data/raw/sample/ 目录。这些是虚构内容，正式使用时删掉即可。
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

OUT_DIR = ROOT / "data" / "raw" / "sample"


def _doc():
    from docx import Document
    return Document()


def write_intro():
    """学校简介：含段落 + 表格。"""
    d = _doc()
    d.add_heading("明德大学简介", level=1)

    d.add_heading("一、学校概况", level=2)
    d.add_paragraph(
        "明德大学创建于1958年，坐落于省会城市东郊文教区，占地面积2100亩，"
        "是一所以工为主，理、工、文、管、艺多学科协调发展的省属重点大学。"
        "学校现有全日制在校本科生21000余人，硕士研究生3600余人，教职工1780人，"
        "其中专任教师1260人，具有高级职称者占48%。"
    )
    d.add_paragraph(
        "学校校训为「明德、笃学、求是、创新」，校庆日为每年10月18日。"
        "学校代码为10425，招生咨询电话为0571-88888000，招生邮箱为 "
        "zhaosheng@mingde.edu.cn。"
    )

    d.add_heading("二、历史沿革", level=2)
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "年份"
    t.rows[0].cells[1].text = "事件"
    for year, event in [
        ("1958年", "前身「明德工业专科学校」成立"),
        ("1978年", "恢复本科招生，更名为明德工学院"),
        ("1995年", "与原明德师范学院合并，定名明德大学"),
        ("2011年", "获批硕士学位授予单位"),
        ("2021年", "入选省高水平大学建设序列"),
    ]:
        row = t.add_row()
        row.cells[0].text = year
        row.cells[1].text = event

    d.add_heading("三、校区与地址", level=2)
    d.add_paragraph(
        "学校现有两个校区：主校区（东湖校区）位于文教区明德路1号，"
        "承担本科一至三年级教学；南湖校区位于南湖大道128号，"
        "主要承担高年级教学、科研与研究生培养。"
    )
    d.add_paragraph(
        "两校区之间开通校车，工作日 7:00—21:00 每30分钟一班，"
        "周末及节假日每60分钟一班，单程约25分钟，教职工凭校园卡免费乘坐。"
    )

    d.add_heading("四、办学条件", level=2)
    d.add_paragraph(
        "图书馆建筑面积3.6万平方米，纸质藏书228万册，电子资源数据库62个。"
        "学校建有国家级实验教学示范中心2个、省级重点实验室9个，"
        "体育馆、游泳馆、标准田径场等体育设施齐全。"
    )

    path = OUT_DIR / "明德大学简介.docx"
    d.save(str(path))
    return path


def write_majors():
    """专业介绍：以表格承载专业核心参数，考验表格解析能力。"""
    d = _doc()
    d.add_heading("本科生专业介绍", level=1)

    d.add_heading("一、专业设置总览", level=2)
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    heads = ["专业名称", "所属学院", "学制", "授予学位", "招生科类"]
    for i, h in enumerate(heads):
        t.rows[0].cells[i].text = h
    for row_data in [
        ("计算机科学与技术", "计算机学院", "4年", "工学学士", "理工类"),
        ("软件工程", "计算机学院", "4年", "工学学士", "理工类"),
        ("人工智能", "计算机学院", "4年", "工学学士", "理工类"),
        ("汉语言文学", "文学院", "4年", "文学学士", "文史类"),
        ("会计学", "商学院", "4年", "管理学学士", "文理兼收"),
        ("视觉传达设计", "艺术学院", "4年", "艺术学学士", "艺术类"),
    ]:
        r = t.add_row()
        for i, v in enumerate(row_data):
            r.cells[i].text = v

    d.add_heading("二、计算机科学与技术专业", level=2)
    d.add_paragraph(
        "本专业培养掌握计算机系统与软件开发基本理论，具备工程实践能力的应用型人才。"
        "核心课程包括：程序设计基础、数据结构与算法、计算机组成原理、操作系统、"
        "计算机网络、数据库系统原理、软件工程、机器学习导论。"
    )
    d.add_paragraph(
        "专业总学分165学分，其中实践教学环节占32%。毕业要求为完成毕业设计并通过答辩。"
        "近年就业率稳定在94%以上，主要去向为软件与信息技术服务业、互联网企业、"
        "金融机构科技部门。"
    )

    d.add_heading("三、人工智能专业", level=2)
    d.add_paragraph(
        "本专业2021年获批设立，2022年开始招生，每年招生规模约60人。"
        "核心课程包括：Python程序设计、机器学习、深度学习、自然语言处理、"
        "计算机视觉、知识图谱、大模型应用开发。"
    )
    d.add_paragraph(
        "专业实验室配备GPU计算服务器，面向本科生开放。"
        "学生在校期间需完成至少一个企业实践项目或学科竞赛作品。"
    )

    d.add_heading("四、转专业与辅修", level=2)
    d.add_paragraph(
        "本科生可在第一学年结束后申请转专业，需满足以下条件："
        "已修课程全部及格、无违纪记录、所申请专业有接收名额。"
        "转专业申请时间为每年6月1日至6月15日，通过学院考核后于秋季学期转入。"
    )
    d.add_paragraph(
        "辅修专业面向二年级及以上学生开放，修满规定学分可获辅修证书，"
        "辅修报名时间为每学期开学第二周。"
    )

    path = OUT_DIR / "专业介绍.docx"
    d.save(str(path))
    return path


def write_canteen():
    """食堂与生活服务：大量表格，且含"第三食堂"这类极易混淆的专有名词。"""
    d = _doc()
    d.add_heading("食堂与生活服务指南", level=1)

    d.add_heading("一、食堂分布", level=2)
    d.add_paragraph(
        "东湖校区现有四个食堂：第一食堂（学生一食堂）、第二食堂、"
        "第三食堂（清真食堂）与教工餐厅。南湖校区设有一食堂和风味餐厅。"
    )

    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(["食堂名称", "位置", "营业时间"]):
        t.rows[0].cells[i].text = h
    for row_data in [
        ("第一食堂", "学生生活区A区一层", "6:30—8:30 / 11:00—13:00 / 17:00—19:00"),
        ("第二食堂", "学生生活区B区二层", "6:30—9:00 / 10:30—13:30 / 16:30—19:30"),
        ("第三食堂", "学生生活区C区一层", "7:00—9:00 / 11:00—13:30 / 17:00—20:00"),
        ("教工餐厅", "行政楼北侧", "11:30—13:00 / 17:30—19:00"),
        ("南湖一食堂", "南湖校区宿舍区", "6:30—8:30 / 11:00—13:00 / 17:00—19:00"),
    ]:
        r = t.add_row()
        for i, v in enumerate(row_data):
            r.cells[i].text = v

    d.add_heading("二、第二食堂特色窗口", level=2)
    t2 = d.add_table(rows=1, cols=3)
    t2.style = "Table Grid"
    for i, h in enumerate(["窗口", "特色菜品", "参考价格"]):
        t2.rows[0].cells[i].text = h
    for row_data in [
        ("2号窗口", "红烧牛肉面、番茄鸡蛋面", "12元 / 9元"),
        ("4号窗口", "瓦罐汤、红烧肉套餐", "15元"),
        ("7号窗口", "麻辣香锅（自选称重）", "约18元"),
        ("9号窗口", "现磨豆浆、葱油饼、小笼包", "3元 / 4元 / 6元"),
    ]:
        r = t2.add_row()
        for i, v in enumerate(row_data):
            r.cells[i].text = v

    d.add_heading("三、第三食堂（清真食堂）", level=2)
    d.add_paragraph(
        "第三食堂为清真食堂，全部窗口严格遵守清真饮食规范，"
        "不供应猪肉及其制品，设有独立的后厨与餐具处理区。"
        "食堂开放时间为早上7:00至9:00、中午11:00至13:30、晚上17:00至20:00，"
        "其中周五中午13:00至13:40为礼拜时间，部分窗口暂停服务。"
    )
    d.add_paragraph(
        "三食堂一楼的「兰州拉面」窗口和二楼的手抓饭窗口价格分别为14元和16元，"
        "是学生评价较高的两个窗口。三食堂支持校园一卡通和手机扫码支付。"
    )

    d.add_heading("四、校园卡与支付", level=2)
    d.add_paragraph(
        "校园一卡通在食堂、超市、图书馆、宿舍门禁通用。"
        "充值方式有两种：一是通过「明德校园通」APP 在线充值，实时到账；"
        "二是到行政楼一楼卡务中心现金充值，工作时间为周一至周五 9:00—16:00。"
        "校园卡遗失请立即通过 APP 挂失，补办工本费20元。"
    )

    d.add_heading("五、图书馆开放时间", level=2)
    t3 = d.add_table(rows=1, cols=2)
    t3.style = "Table Grid"
    t3.rows[0].cells[0].text = "区域"
    t3.rows[0].cells[1].text = "开放时间"
    for row_data in [
        ("中文图书借阅区", "周一至周日 8:00—22:00"),
        ("期刊阅览室", "周一至周五 8:30—21:00"),
        ("电子阅览室", "周一至周日 9:00—21:30"),
        ("自习区", "周一至周日 7:00—23:00"),
        ("研讨间（需预约）", "周一至周日 9:00—21:00"),
    ]:
        r = t3.add_row()
        for i, v in enumerate(row_data):
            r.cells[i].text = v
    d.add_paragraph("法定节假日开放时间另行通知，寒暑假期间开馆时间调整为 9:00—16:00。")

    d.add_heading("六、宿舍与后勤报修", level=2)
    d.add_paragraph(
        "本科生宿舍为四人间，配备独立卫浴、空调和网络接口。"
        "宿舍供电时间为 6:00—23:30，考试周延长至 24:00。"
        "后勤报修电话为 0571-88888123，也可通过「明德校园通」APP 提交报修单，"
        "一般故障24小时内响应。"
    )

    path = OUT_DIR / "食堂与生活服务.docx"
    d.save(str(path))
    return path


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    files = [write_intro(), write_majors(), write_canteen()]
    print("已生成演示文档：")
    for f in files:
        print(f"  {f.relative_to(ROOT)}  ({f.stat().st_size // 1024} KB)")
    print("\n下一步：python -m app.cli build")


if __name__ == "__main__":
    main()
